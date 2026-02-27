#!/usr/bin/env python3
"""
Excel Reader - Excelファイル読み取りスクリプト

対応形式: .xlsx (Office Open XML)
依存ライブラリ: openpyxl
"""
import argparse
import io
import csv
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("エラー: openpyxlがインストールされていません。pip install openpyxl を実行してください。", file=sys.stderr)
    sys.exit(1)


def cmd_info(args: argparse.Namespace) -> None:
    """シート一覧と基本情報を表示"""
    wb = openpyxl.load_workbook(args.file_path, read_only=True, data_only=True)
    print(f"ファイル: {args.file_path}")
    print(f"シート数: {len(wb.sheetnames)}")
    print()
    for name in wb.sheetnames:
        ws = wb[name]
        print(f"  シート: {name}")
        print(f"    行数: {ws.max_row}, 列数: {ws.max_column}")
    wb.close()


def cmd_csv(args: argparse.Namespace) -> None:
    """シートデータをCSV形式で出力"""
    wb = openpyxl.load_workbook(args.file_path, read_only=True, data_only=True)

    sheet_name = args.sheet
    if not sheet_name:
        sheet_name = wb.sheetnames[0]

    if sheet_name not in wb.sheetnames:
        print(f"エラー: シート '{sheet_name}' が見つかりません。", file=sys.stderr)
        print(f"利用可能なシート: {', '.join(wb.sheetnames)}", file=sys.stderr)
        wb.close()
        sys.exit(1)

    ws = wb[sheet_name]
    start_row = args.start_row or 1
    end_row = args.end_row or ws.max_row

    output = io.StringIO()
    writer = csv.writer(output)

    for row in ws.iter_rows(min_row=start_row, max_row=end_row, values_only=True):
        writer.writerow([str(cell) if cell is not None else "" for cell in row])

    print(output.getvalue(), end="")
    wb.close()


def cmd_search(args: argparse.Namespace) -> None:
    """ワークブック全体からキーワード検索"""
    wb = openpyxl.load_workbook(args.file_path, read_only=True, data_only=True)
    query = args.query
    case_sensitive = args.case_sensitive
    results = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            for col_idx, cell in enumerate(row, start=1):
                if cell is None:
                    continue
                cell_str = str(cell)
                if case_sensitive:
                    match = query in cell_str
                else:
                    match = query.lower() in cell_str.lower()
                if match:
                    results.append({
                        "sheet": sheet_name,
                        "row": row_idx,
                        "col": col_idx,
                        "value": cell_str,
                    })

    if not results:
        print(f"'{query}' は見つかりませんでした。")
    else:
        print(f"検索結果: {len(results)}件")
        print()
        for r in results:
            print(f"  シート: {r['sheet']}, 行: {r['row']}, 列: {r['col']}")
            value_preview = r["value"][:100]
            if len(r["value"]) > 100:
                value_preview += "..."
            print(f"    値: {value_preview}")

    wb.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="Excel Reader")
    subparsers = parser.add_subparsers(dest="command", required=True)

    # info
    p_info = subparsers.add_parser("info", help="シート一覧と基本情報を表示")
    p_info.add_argument("file_path", help="Excelファイルパス")

    # csv
    p_csv = subparsers.add_parser("csv", help="シートデータをCSV形式で出力")
    p_csv.add_argument("file_path", help="Excelファイルパス")
    p_csv.add_argument("--sheet", default="", help="シート名（省略時は最初のシート）")
    p_csv.add_argument("--start-row", type=int, default=None, help="開始行")
    p_csv.add_argument("--end-row", type=int, default=None, help="終了行")

    # search
    p_search = subparsers.add_parser("search", help="ワークブック全体からキーワード検索")
    p_search.add_argument("file_path", help="Excelファイルパス")
    p_search.add_argument("query", help="検索キーワード")
    p_search.add_argument("--case-sensitive", action="store_true", help="大文字小文字を区別")

    args = parser.parse_args()

    if args.command == "info":
        cmd_info(args)
    elif args.command == "csv":
        cmd_csv(args)
    elif args.command == "search":
        cmd_search(args)


if __name__ == "__main__":
    main()
