#!/usr/bin/env python3
"""
DOCX Reader - Word文書解析スクリプト (AI最適化版)

機能:
  - メタデータ抽出（作成者、日付等）
  - 見出し構造（H1-H9）の抽出
  - 段落テキスト抽出（範囲指定可能）
  - 表の構造化出力
  - 見出しジャンプ（セクション単位での表示）
  - キーワード検索（段落＋表を横断）
  - AI向けコンパクト出力 / JSON出力 / 人間向けテキスト出力

使用例:
  python3 read_docx.py document.docx                          # AI向け出力(デフォルト)
  python3 read_docx.py document.docx --format human            # 人間向け出力
  python3 read_docx.py document.docx --format json              # JSON出力(空値省略)
  python3 read_docx.py document.docx --paragraphs 1-50
  python3 read_docx.py document.docx --heading "第3章"
  python3 read_docx.py document.docx --search "キーワード"
  python3 read_docx.py document.docx --summary                  # 見出し構造のみ
"""

from __future__ import annotations

import argparse
import json
import subprocess
import sys
import unicodedata
from dataclasses import dataclass, field, asdict
from pathlib import Path
from typing import Optional


def ensure_python_docx() -> None:
    """python-docx がなければインストールする。"""
    try:
        import docx  # noqa: F401
    except ImportError:
        print("python-docx をインストール中...", file=sys.stderr)
        subprocess.check_call(
            [sys.executable, "-m", "pip", "install", "python-docx",
             "--break-system-packages", "-q"],
        )


ensure_python_docx()

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402


# ---------------------------------------------------------------------------
# テキスト正規化
# ---------------------------------------------------------------------------

def _normalize_text(text: str) -> str:
    """Unicode NFC正規化 + 制御文字除去。"""
    if not text:
        return ""
    text = unicodedata.normalize("NFC", text)
    result = []
    for ch in text:
        code = ord(ch)
        if ch in ("\n", "\r", "\t"):
            result.append(ch)
        elif code < 0x20 or code == 0x7F or (0x80 <= code <= 0x9F):
            continue
        elif code in (0x200B, 0x200C, 0x200D, 0x2060, 0xFEFF):
            continue
        else:
            result.append(ch)
    return "".join(result)


# ---------------------------------------------------------------------------
# データクラス
# ---------------------------------------------------------------------------

@dataclass
class HeadingInfo:
    level: int
    text: str
    para_index: int
    char_count: int = 0


@dataclass
class TableInfo:
    index: int
    rows: int
    cols: int
    near_para: int
    cells: list[list[str]] = field(default_factory=list)
    has_header: bool = False


@dataclass
class ParagraphInfo:
    index: int
    text: str
    style: str = ""
    heading_level: int = 0
    bold_texts: list[str] = field(default_factory=list)


@dataclass
class MetadataInfo:
    author: str = ""
    last_modified_by: str = ""
    created: str = ""
    modified: str = ""
    revision: Optional[int] = None
    title: str = ""
    subject: str = ""
    category: str = ""


@dataclass
class DocumentInfo:
    filename: str = ""
    total_paragraphs: int = 0
    total_characters: int = 0
    tables_count: int = 0
    metadata: MetadataInfo = field(default_factory=MetadataInfo)
    headings: list[HeadingInfo] = field(default_factory=list)
    tables: list[TableInfo] = field(default_factory=list)
    paragraphs: list[ParagraphInfo] = field(default_factory=list)


# ---------------------------------------------------------------------------
# パーサー
# ---------------------------------------------------------------------------

def _safe_str(value: object) -> str:
    return str(value) if value else ""


def _get_heading_level(paragraph) -> int:
    """段落の見出しレベルを取得する (0=見出しでない, 1-9=Heading level)。"""
    style_name = paragraph.style.name if paragraph.style else ""
    if style_name.startswith("Heading"):
        try:
            return int(style_name.replace("Heading", "").strip())
        except ValueError:
            return 0
    # Japanese style names
    if "見出し" in style_name:
        for ch in style_name:
            if ch.isdigit():
                return int(ch)
        return 1
    return 0


def _detect_table_header(table) -> bool:
    """テーブルの1行目がヘッダー行か自動判定する。"""
    rows = table.rows
    if len(rows) < 2:
        return False
    first_texts = [cell.text.strip() for cell in rows[0].cells]
    if all(0 < len(t) < 30 for t in first_texts):
        return True
    return False


def parse_docx(filepath: str) -> DocumentInfo:
    doc = Document(filepath)
    props = doc.core_properties

    info = DocumentInfo(
        filename=Path(filepath).name,
        metadata=MetadataInfo(
            author=_safe_str(props.author),
            last_modified_by=_safe_str(props.last_modified_by),
            created=_safe_str(props.created),
            modified=_safe_str(props.modified),
            revision=props.revision,
            title=_safe_str(props.title),
            subject=_safe_str(props.subject),
            category=_safe_str(props.category),
        ),
    )

    # Parse paragraphs
    total_chars = 0
    for i, para in enumerate(doc.paragraphs):
        text = _normalize_text(para.text)
        heading_level = _get_heading_level(para)

        # Collect bold runs
        bold_texts = []
        for run in para.runs:
            if run.bold and run.text.strip():
                bold_texts.append(run.text.strip())

        pi = ParagraphInfo(
            index=i + 1,
            text=text,
            style=para.style.name if para.style else "",
            heading_level=heading_level,
            bold_texts=bold_texts,
        )
        info.paragraphs.append(pi)

        if heading_level > 0:
            # Count characters until next heading
            info.headings.append(HeadingInfo(
                level=heading_level,
                text=text,
                para_index=i + 1,
            ))

        total_chars += len(text)

    info.total_paragraphs = len(info.paragraphs)
    info.total_characters = total_chars

    # Compute heading char_count (chars from this heading to next heading)
    for idx, heading in enumerate(info.headings):
        start = heading.para_index
        if idx + 1 < len(info.headings):
            end = info.headings[idx + 1].para_index
        else:
            end = info.total_paragraphs + 1
        char_count = sum(
            len(p.text) for p in info.paragraphs
            if start <= p.index < end
        )
        heading.char_count = char_count

    # Parse tables
    para_count = 0
    for element in doc.element.body:
        tag = element.tag.split("}")[-1] if "}" in element.tag else element.tag
        if tag == "p":
            para_count += 1
        elif tag == "tbl":
            # Find corresponding table
            for ti, table in enumerate(doc.tables):
                if table._element is element:
                    rows = table.rows
                    cells = []
                    for row in rows:
                        cells.append([_normalize_text(cell.text) for cell in row.cells])

                    t_info = TableInfo(
                        index=ti + 1,
                        rows=len(rows),
                        cols=len(rows[0].cells) if rows else 0,
                        near_para=para_count,
                        cells=cells,
                        has_header=_detect_table_header(table),
                    )
                    info.tables.append(t_info)
                    break

    info.tables_count = len(info.tables)
    return info


# ---------------------------------------------------------------------------
# 範囲パーサー
# ---------------------------------------------------------------------------

def parse_range(spec: str, total: int) -> list[int]:
    result: set[int] = set()
    for part in spec.split(","):
        part = part.strip()
        if "-" in part:
            s, e = part.split("-", 1)
            result.update(range(max(1, int(s)), min(total, int(e)) + 1))
        else:
            n = int(part)
            if 1 <= n <= total:
                result.add(n)
    return sorted(result)


# ---------------------------------------------------------------------------
# キーワード検索
# ---------------------------------------------------------------------------

@dataclass
class SearchHit:
    location: str
    para_index: int
    text: str
    context: str


def search_document(info: DocumentInfo, keyword: str) -> list[SearchHit]:
    kw = keyword.lower()
    hits: list[SearchHit] = []

    for para in info.paragraphs:
        if kw in para.text.lower():
            hits.append(SearchHit(
                location=f"P{para.index}",
                para_index=para.index,
                text=para.text,
                context=_extract_context(para.text, kw),
            ))

    for table in info.tables:
        for ri, row in enumerate(table.cells):
            for ci, cell_text in enumerate(row):
                if kw in cell_text.lower():
                    hits.append(SearchHit(
                        location=f"表{table.index}[行{ri + 1},列{ci + 1}]",
                        para_index=table.near_para,
                        text=cell_text,
                        context=cell_text[:80],
                    ))

    return hits


def _extract_context(text: str, kw: str, window: int = 40) -> str:
    idx = text.lower().find(kw)
    if idx == -1:
        return text[:80]
    s = max(0, idx - window)
    e = min(len(text), idx + len(kw) + window)
    pre = "…" if s > 0 else ""
    suf = "…" if e < len(text) else ""
    return f"{pre}{text[s:e]}{suf}"


# ---------------------------------------------------------------------------
# 見出しジャンプ
# ---------------------------------------------------------------------------

def find_heading_section(info: DocumentInfo, heading_query: str) -> Optional[tuple[HeadingInfo, list[int]]]:
    """見出しを部分一致で検索し、そのセクションの段落範囲を返す。"""
    q = heading_query.lower()
    for idx, heading in enumerate(info.headings):
        if q in heading.text.lower():
            start = heading.para_index
            if idx + 1 < len(info.headings):
                end = info.headings[idx + 1].para_index - 1
            else:
                end = info.total_paragraphs
            return heading, list(range(start, end + 1))
    return None


# ---------------------------------------------------------------------------
# AI向けコンパクト出力
# ---------------------------------------------------------------------------

def format_ai(info: DocumentInfo, para_filter: Optional[list[int]] = None,
              max_paragraphs: int = 50) -> str:
    lines: list[str] = []
    m = info.metadata

    # ヘッダー
    meta_parts: list[str] = [f"{info.total_paragraphs}段落", f"{info.total_characters:,}字"]
    if info.filename:
        meta_parts.insert(0, info.filename)
    if info.tables_count:
        meta_parts.append(f"表{info.tables_count}個")
    lines.append(f"[DOCX] {' | '.join(meta_parts)}")

    # 目次
    if info.headings:
        toc_parts: list[str] = []
        for h in info.headings:
            toc_parts.append(f"H{h.level}:{h.text}(P{h.para_index})")
        lines.append(f"[目次] {' / '.join(toc_parts)}")

    lines.append("")

    # Determine which paragraphs to show
    if para_filter:
        target_indices = set(para_filter)
    else:
        target_indices = set(range(1, min(info.total_paragraphs + 1, max_paragraphs + 1)))

    # Group paragraphs by section
    current_heading: Optional[HeadingInfo] = None
    section_header_shown = False

    for para in info.paragraphs:
        if para.index not in target_indices:
            # Check if this is a heading that defines a section
            if para.heading_level > 0:
                current_heading = HeadingInfo(
                    level=para.heading_level,
                    text=para.text,
                    para_index=para.index,
                )
                section_header_shown = False
            continue

        # Show section header
        if para.heading_level > 0:
            # Find char_count for this heading
            char_count = ""
            for h in info.headings:
                if h.para_index == para.index:
                    char_count = f", 約{h.char_count:,}字" if h.char_count else ""
                    break
            lines.append(f"== H{para.heading_level}: {para.text} (P{para.index}{char_count})")
            current_heading = HeadingInfo(
                level=para.heading_level,
                text=para.text,
                para_index=para.index,
            )
            section_header_shown = True
        elif para.text.strip():
            if current_heading and not section_header_shown:
                lines.append(f"== H{current_heading.level}: {current_heading.text} (P{current_heading.para_index})")
                section_header_shown = True

            bold_note = ""
            if para.bold_texts:
                bold_note = f" (強調: {'、'.join(para.bold_texts)})"
            lines.append(f"  {para.text}{bold_note}")

        # Show tables near this paragraph
        for table in info.tables:
            if table.near_para == para.index:
                lines.append(f"  [表{table.index}] {table.rows}行x{table.cols}列")
                for ri, row in enumerate(table.cells):
                    prefix = "H" if (ri == 0 and table.has_header) else str(ri + 1)
                    lines.append(f"    {prefix}| {' | '.join(row)}")

    if not para_filter and info.total_paragraphs > max_paragraphs:
        lines.append("")
        remaining = info.total_paragraphs - max_paragraphs
        lines.append(f"--- 残り{remaining}段落。取得: --paragraphs \"{max_paragraphs + 1}-{info.total_paragraphs}\" ---")

    return "\n".join(lines).rstrip()


def format_search_ai(hits: list[SearchHit], keyword: str) -> str:
    if not hits:
        return f"検索: '{keyword}' → 該当なし"
    lines = [f"検索: '{keyword}' → {len(hits)}件"]
    for h in hits:
        lines.append(f"  {h.location} {h.context}")
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# 人間向け出力
# ---------------------------------------------------------------------------

def format_human(info: DocumentInfo, para_filter: Optional[list[int]] = None,
                 max_paragraphs: int = 50) -> str:
    lines: list[str] = []
    lines.append("=" * 60)
    lines.append("【Word文書情報】")
    lines.append(f"  ファイル名: {info.filename}")
    lines.append(f"  段落数: {info.total_paragraphs}")
    lines.append(f"  文字数: {info.total_characters:,}")
    lines.append(f"  表の数: {info.tables_count}")
    m = info.metadata
    for label, val in [("タイトル", m.title), ("作成者", m.author),
                       ("最終更新者", m.last_modified_by),
                       ("作成日", m.created), ("更新日", m.modified)]:
        if val:
            lines.append(f"  {label}: {val}")
    lines.append("=" * 60)

    if info.headings:
        lines.append("\n【見出し構造】")
        for h in info.headings:
            indent = "  " * h.level
            lines.append(f"  {indent}H{h.level}: {h.text} (P{h.para_index})")

    target_indices = set(para_filter) if para_filter else set(
        range(1, min(info.total_paragraphs + 1, max_paragraphs + 1)))

    for para in info.paragraphs:
        if para.index not in target_indices:
            continue
        if para.heading_level > 0:
            lines.append(f"\n--- H{para.heading_level}: {para.text} (P{para.index}) ---")
        elif para.text.strip():
            lines.append(f"  {para.text}")

        for table in info.tables:
            if table.near_para == para.index:
                hdr = " (ヘッダー付き)" if table.has_header else ""
                lines.append(f"\n  [表{table.index}] {table.rows}行 x {table.cols}列{hdr}")
                for ri, row in enumerate(table.cells):
                    label = "H" if (ri == 0 and table.has_header) else str(ri + 1)
                    lines.append(f"    {label}: {row}")

    lines.append("\n" + "=" * 60)
    return "\n".join(lines)


def format_search_human(hits: list[SearchHit], keyword: str) -> str:
    if not hits:
        return f"キーワード '{keyword}' は見つかりませんでした。"
    lines = [f"検索結果: '{keyword}' ({len(hits)}件)", "=" * 40]
    for i, h in enumerate(hits, 1):
        lines.append(f"  [{i}] {h.location}: {h.context}")
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# JSON出力 (空値省略)
# ---------------------------------------------------------------------------

def _compact_dict(d: dict) -> dict:
    out = {}
    for k, v in d.items():
        if v is None or v == "" or v == [] or v is False or v == 0:
            continue
        if isinstance(v, dict):
            v = _compact_dict(v)
            if not v:
                continue
        elif isinstance(v, list):
            v = [_compact_dict(i) if isinstance(i, dict) else i for i in v]
            v = [i for i in v if i]
            if not v:
                continue
        out[k] = v
    return out


def format_json(info: DocumentInfo, para_filter: Optional[list[int]] = None,
                max_paragraphs: int = 50) -> str:
    d = asdict(info)
    if para_filter:
        d["paragraphs"] = [p for p in d["paragraphs"] if p["index"] in para_filter]
    else:
        d["paragraphs"] = d["paragraphs"][:max_paragraphs]
    return json.dumps(_compact_dict(d), ensure_ascii=False, indent=2)


# ---------------------------------------------------------------------------
# 概要のみ出力
# ---------------------------------------------------------------------------

def format_summary(info: DocumentInfo) -> str:
    lines: list[str] = []
    meta_parts = [info.filename, f"{info.total_paragraphs}段落",
                  f"{info.total_characters:,}字"]
    if info.tables_count:
        meta_parts.append(f"表{info.tables_count}個")
    lines.append(f"[DOCX] {' | '.join(meta_parts)}")

    if info.headings:
        lines.append("")
        for h in info.headings:
            indent = "  " * (h.level - 1)
            char_label = f" ({h.char_count:,}字)" if h.char_count else ""
            lines.append(f"  {indent}H{h.level}: {h.text} (P{h.para_index}){char_label}")

    if info.tables:
        lines.append("")
        for t in info.tables:
            lines.append(f"  [表{t.index}] {t.rows}行x{t.cols}列 (P{t.near_para}付近)")

    return "\n".join(lines)


# ---------------------------------------------------------------------------
# メイン
# ---------------------------------------------------------------------------

def main() -> None:
    parser = argparse.ArgumentParser(description="DOCX Reader - Word文書解析ツール")
    parser.add_argument("file", help="Wordファイルのパス")
    parser.add_argument("--paragraphs", help="段落範囲 (例: 1-50, 20,30,40)")
    parser.add_argument("--heading", help="見出しジャンプ（部分一致でセクション表示）")
    parser.add_argument("--search", help="キーワード検索")
    parser.add_argument("--format", choices=["ai", "human", "json"], default="ai",
                        help="出力形式 (デフォルト: ai)")
    parser.add_argument("--summary", action="store_true", help="見出し構造のみ出力")
    parser.add_argument("--max-paragraphs", type=int, default=50,
                        help="最大段落数 (デフォルト: 50)")

    args = parser.parse_args()
    filepath = Path(args.file)

    if not filepath.exists():
        print(f"エラー: ファイルが見つかりません: {filepath}", file=sys.stderr)
        sys.exit(1)

    ext = filepath.suffix.lower()
    if ext == ".doc":
        print("エラー: .doc（旧形式）は非対応です。\n"
              ".docx（Office Open XML）形式に変換してください。\n\n"
              "対処方法:\n"
              "1. Wordで .docx 形式に変換して再アップロード\n"
              "2. LibreOffice で .docx 形式に変換", file=sys.stderr)
        sys.exit(1)

    if ext not in (".docx",):
        print(f"エラー: Wordファイルではありません: {filepath}", file=sys.stderr)
        sys.exit(1)

    info = parse_docx(str(filepath))

    para_filter: Optional[list[int]] = None

    if args.paragraphs:
        para_filter = parse_range(args.paragraphs, info.total_paragraphs)

    if args.heading:
        result = find_heading_section(info, args.heading)
        if result:
            heading, para_range = result
            para_filter = para_range
        else:
            print(f"見出し '{args.heading}' が見つかりません。", file=sys.stderr)
            print("", file=sys.stderr)
            print("利用可能な見出し:", file=sys.stderr)
            for h in info.headings:
                print(f"  H{h.level}: {h.text}", file=sys.stderr)
            sys.exit(1)

    if args.summary:
        if args.format == "json":
            # Summary as JSON: exclude paragraphs
            d = asdict(info)
            del d["paragraphs"]
            print(json.dumps(_compact_dict(d), ensure_ascii=False, indent=2))
        elif args.format == "human":
            print(format_human(info, para_filter=[], max_paragraphs=0))
        else:
            print(format_summary(info))
        return

    if args.search:
        hits = search_document(info, args.search)
        if args.format == "json":
            print(json.dumps([_compact_dict(asdict(h)) for h in hits],
                             ensure_ascii=False, indent=2))
        elif args.format == "human":
            print(format_search_human(hits, args.search))
        else:
            print(format_search_ai(hits, args.search))
        return

    if args.format == "json":
        print(format_json(info, para_filter, args.max_paragraphs))
    elif args.format == "human":
        print(format_human(info, para_filter, args.max_paragraphs))
    else:
        print(format_ai(info, para_filter, args.max_paragraphs))


if __name__ == "__main__":
    main()
